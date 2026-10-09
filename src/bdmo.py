from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin, urlparse
import hashlib
import re
import shutil
import tempfile
import zipfile

import numpy as np
import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from .rules import FEATURE_RULES


COLUMN_CANDIDATES = {
    'oktmo': ['октмо', 'oktmo', 'код октмо'],
    'year': ['год', 'year'],
    'value': ['indicator_value', 'значение', 'value', 'значение показателя'],
    'indicator_name': ['показатель', 'наименование показателя', 'название показателя', 'indicator', 'indicator_name'],
    'unit': ['indicator_unit', 'единица измерения', 'ед. изм.', 'unit'],
    'municipality_name': ['municipality_name', 'муниципальное образование', 'наименование муниципального образования', 'мо', 'municipality'],
    'municipality_type': ['municipality_type', 'mun_type', 'тип муниципального образования', 'тип мо', 'municipality type'],
    'region_name': ['region_name', 'регион', 'субъект российской федерации', 'наименование региона', 'region'],
    'period': ['indicator_period', 'период', 'period'],
}

TOTAL_MARKERS = {
    '', 'всего', 'итого', 'все', 'все организации', 'все формы собственности',
    'всего по обследуемым видам экономической деятельности'
}


def normalize_text(value):
    if value is None or pd.isna(value):
        return ''
    return re.sub(r'\s+', ' ', str(value).replace('\xa0', ' ')).strip()


def text_key(value):
    return normalize_text(value).casefold()


def find_column(columns, candidates):
    keyed = {text_key(column): column for column in columns}

    for candidate in candidates:
        if text_key(candidate) in keyed:
            return keyed[text_key(candidate)]

    for candidate in candidates:
        key = text_key(candidate)
        for normalized, original in keyed.items():
            if key in normalized:
                return original

    return None


def normalize_oktmo(value):
    # Do not invent leading zeros or strip arbitrary characters from identifiers.
    text = normalize_text(value)
    text = re.sub(r'\.0$', '', text)
    return text if re.fullmatch(r'\d{8}|\d{11}', text) else None


def numeric_value(series):
    text = (
        series.astype('string')
        .str.replace('\xa0', '', regex=False)
        .str.replace(' ', '', regex=False)
        .str.replace(',', '.', regex=False)
    )
    text = text.replace({
        '': pd.NA, '-': pd.NA, '—': pd.NA, '...': pd.NA, '…': pd.NA,
        'н/д': pd.NA, 'нет данных': pd.NA,
    })
    return pd.to_numeric(text, errors='coerce').replace([np.inf, -np.inf], np.nan)


def make_session():
    session = requests.Session()
    session.headers.update({'User-Agent': 'Mozilla/5.0 BDMO-research/1.0'})
    return session


def discover_source(source_page, session, timeout=120):
    response = session.get(source_page, timeout=timeout)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, 'html.parser')

    rows = []
    for anchor in soup.find_all('a', href=True):
        href = urljoin(source_page, anchor['href'])
        if 'storage.yandexcloud.net' not in href:
            continue

        context = normalize_text(anchor.parent.get_text(' ', strip=True))
        match = re.search(r'(?<!\d)(\d{1,2})\.\s*(.+?)(?:\s*\(|$)', context)
        if not match:
            continue

        link_text = text_key(anchor.get_text(' ', strip=True))
        if 'целиком' in link_text:
            mode = 'section'
        elif 'разбивке' in link_text:
            mode = 'by_indicator'
        else:
            mode = 'other'

        rows.append({
            'section': int(match.group(1)),
            'section_title': normalize_text(match.group(2)),
            'mode': mode,
            'url': href,
        })

    section_links = pd.DataFrame(rows).drop_duplicates()
    if section_links.empty:
        raise RuntimeError('На странице БДМО не найдены ссылки на разделы.')

    sheet_links = [
        urljoin(source_page, anchor['href'])
        for anchor in soup.find_all('a', href=True)
        if 'docs.google.com/spreadsheets' in anchor['href']
    ]
    if not sheet_links:
        raise RuntimeError('На странице БДМО не найдена ссылка на справочник показателей.')

    return section_links.sort_values(['section', 'mode']).reset_index(drop=True), sheet_links[0]


def _download(url, directory, session, timeout=120):
    response = session.get(url, stream=True, timeout=timeout)
    response.raise_for_status()

    suffix = Path(urlparse(response.url).path).suffix
    if not suffix:
        suffix = Path(urlparse(url).path).suffix
    if not suffix:
        content_type = response.headers.get('content-type', '').casefold()
        suffix = '.zip' if 'zip' in content_type else '.bin'

    name = hashlib.sha256(url.encode('utf-8')).hexdigest()[:12] + suffix
    destination = Path(directory) / name

    with destination.open('wb') as file:
        for chunk in response.iter_content(chunk_size=1024 * 1024):
            if chunk:
                file.write(chunk)

    return destination


def _sheet_export_url(sheet_url):
    match = re.search(r'/d/([^/]+)', sheet_url)
    if not match:
        raise RuntimeError('Не удалось определить id Google Sheets.')
    sheet_id = match.group(1)
    return f'https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx'


def _detect_header_row(worksheet, max_rows=20):
    best_row = 1
    best_score = -1

    for row_number in range(1, min(worksheet.max_row, max_rows) + 1):
        values = [
            text_key(worksheet.cell(row_number, column).value)
            for column in range(1, worksheet.max_column + 1)
        ]
        nonempty = sum(bool(value) for value in values)
        semantic = sum(
            any(token in value for token in ['раздел', 'показател', 'скач', 'код', 'наимен'])
            for value in values
        )
        score = nonempty + 5 * semantic
        if score > best_score:
            best_row = row_number
            best_score = score

    return best_row


def _formula_url(value):
    if not isinstance(value, str):
        return None
    match = re.search(r'HYPERLINK\(["\']([^"\']+)["\']', value, flags=re.IGNORECASE)
    return match.group(1) if match else None


def _worksheet_records(worksheet):
    header_row = _detect_header_row(worksheet)
    headers = []

    for column in range(1, worksheet.max_column + 1):
        value = normalize_text(worksheet.cell(header_row, column).value)
        headers.append(value or f'column_{column}')

    records = []
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        row = {}
        urls = []

        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(row_number, column)
            row[header] = cell.value
            if cell.hyperlink:
                urls.append(cell.hyperlink.target)
            formula_link = _formula_url(cell.value)
            if formula_link:
                urls.append(formula_link)

        if not any(normalize_text(value) for value in row.values()):
            continue

        row['_urls'] = list(dict.fromkeys(urls))
        records.append(row)

    return pd.DataFrame(records), header_row


def _parse_section(value):
    if value is None or pd.isna(value):
        return np.nan

    if isinstance(value, (int, np.integer)):
        return int(value)

    if isinstance(value, (float, np.floating)) and np.isfinite(value):
        if float(value).is_integer():
            return int(value)

    text = normalize_text(value)
    if not text:
        return np.nan

    match = re.search(r'(?<!\d)(\d{1,3})(?:\.0+)?(?!\d)', text)
    return int(match.group(1)) if match else np.nan


def load_indicator_reference(sheet_url, session, timeout=120):
    with tempfile.TemporaryDirectory() as temporary:
        path = _download(_sheet_export_url(sheet_url), temporary, session, timeout)
        content = Path(path).read_bytes()

        if not content.startswith(b'PK'):
            raise RuntimeError(
                'Справочник показателей скачался не как XLSX. '
                'Возможно, Google Sheets вернул HTML-страницу вместо файла.'
            )

        workbook = load_workbook(
            BytesIO(content),
            read_only=False,
            data_only=False,
        )

        sheet_name = next(
            (name for name in workbook.sheetnames if 'показател' in text_key(name)),
            None,
        )
        if sheet_name is None:
            raise RuntimeError(f'В справочнике не найден лист с показателями: {workbook.sheetnames}')

        reference, header_row = _worksheet_records(workbook[sheet_name])

    # В актуальном справочнике БДМО номер раздела хранится отдельно
    # от его текстового названия: «Код раздела» и «Раздел».
    section_column = find_column(
        reference.columns,
        ['код раздела', 'номер раздела', '№ раздела', 'раздел №', 'section code'],
    )
    if section_column is None:
        section_column = find_column(reference.columns, ['section', 'раздел'])

    indicator_column = find_column(
        reference.columns,
        ['название показателя', 'наименование показателя', 'indicator_name', 'indicator'],
    )
    if indicator_column is None:
        indicator_column = find_column(reference.columns, ['показатель'])

    indicator_code_column = find_column(
        reference.columns,
        ['код показателя', 'indicator code', 'indicator_code'],
    )
    unit_column = find_column(
        reference.columns,
        ['единица измерения', 'ед. изм.', 'unit'],
    )
    section_title_column = find_column(
        reference.columns,
        ['раздел'],
    )

    if section_column is None or indicator_column is None:
        raise RuntimeError(
            'Не удалось распознать обязательные колонки справочника. '
            f'Нужны код раздела и название показателя. Колонки: {list(reference.columns)}'
        )

    reference['_section'] = reference[section_column].map(_parse_section).ffill()
    reference['_indicator_name'] = reference[indicator_column].map(normalize_text)

    if indicator_code_column is not None:
        reference['_indicator_code'] = reference[indicator_code_column].map(normalize_text)
    else:
        reference['_indicator_code'] = ''

    if unit_column is not None:
        reference['_unit'] = reference[unit_column].map(normalize_text)
    else:
        reference['_unit'] = ''

    reference = reference[reference['_indicator_name'].ne('')].copy()

    sections = sorted(reference['_section'].dropna().astype(int).unique().tolist())
    if not sections:
        examples = reference[section_column].dropna().head(10).map(normalize_text).tolist()
        raise RuntimeError(
            'Не удалось прочитать коды разделов из справочника. '
            f'Использована колонка {section_column!r}; примеры: {examples}. '
            f'Колонки: {list(reference.columns)}'
        )

    info = {
        'sheet_name': sheet_name,
        'header_row': header_row,
        'section_column': section_column,
        'section_title_column': section_title_column,
        'indicator_column': indicator_column,
        'indicator_code_column': indicator_code_column,
        'unit_column': unit_column,
        'sections': sections,
        'n_indicators': len(reference),
        'columns': list(reference.columns),
    }
    return reference, info


def _rule_score(indicator_name, section, rule):
    if section not in rule['sections']:
        return None

    text = text_key(indicator_name)
    include_hits = sum(token.casefold() in text for token in rule['include_any'])
    if include_hits == 0:
        return None

    required = rule.get('require_any', [])
    if required and not any(token.casefold() in text for token in required):
        return None

    if any(token.casefold() in text for token in rule['exclude_any']):
        return None

    prefer_hits = sum(token.casefold() in text for token in rule['prefer_any'])
    return 100 * include_hits + 20 * prefer_hits


def shortlist_indicators(reference, rules=FEATURE_RULES):
    rows = []

    for rule in rules:
        local = []
        for _, row in reference.iterrows():
            section = row['_section']
            if pd.isna(section):
                continue

            indicator_name = row['_indicator_name']
            score = _rule_score(indicator_name, int(section), rule)
            if score is None:
                continue

            urls = row.get('_urls', [])
            local.append({
                'feature_key': rule['feature_key'],
                'view': rule['view'],
                'section': int(section),
                'indicator_code': row.get('_indicator_code', ''),
                'indicator_name': indicator_name,
                'unit': row.get('_unit', ''),
                'match_score': score,
                'urls': urls,
            })

        local.sort(key=lambda item: (-item['match_score'], item['indicator_name']))
        rows.extend(local[:rule['max_indicators']])

    candidates = pd.DataFrame(rows)
    if candidates.empty:
        sections = sorted(
            reference['_section'].dropna().astype(int).unique().tolist()
        )
        examples = reference['_indicator_name'].dropna().head(10).tolist()
        raise RuntimeError(
            'Автоматический shortlist показателей пуст. '
            f'Распознанные разделы: {sections}. '
            f'Примеры показателей: {examples}'
        )
    return candidates


def _choose_url(urls):
    if not isinstance(urls, list):
        return None

    usable = [url for url in urls if isinstance(url, str) and url.startswith('http')]
    for suffix in ['.parquet', '.zip', '.csv', '.xlsx']:
        for url in usable:
            if suffix in url.split('?')[0].casefold():
                return url
    return usable[0] if usable else None


def attach_download_urls(candidates, section_links):
    result = candidates.copy()
    result['download_url'] = result['urls'].map(_choose_url)

    fallback = {}
    for section in sorted(result['section'].unique()):
        part = section_links[section_links['section'].eq(section)]
        if part.empty:
            continue
        whole = part[part['mode'].eq('section')]
        split = part[part['mode'].eq('by_indicator')]
        if not whole.empty:
            fallback[section] = whole.iloc[0]['url']
        elif not split.empty:
            fallback[section] = split.iloc[0]['url']

    result['download_url'] = result.apply(
        lambda row: row['download_url'] or fallback.get(row['section']), axis=1
    )

    missing = result[result['download_url'].isna()]
    if len(missing):
        keys = missing['feature_key'].tolist()
        raise RuntimeError(f'Не найдены ссылки на данные для: {keys}')

    return result.drop(columns=['urls'])


def _extract(path, destination):
    destination.mkdir(parents=True, exist_ok=True)
    if path.suffix.casefold() == '.zip':
        with zipfile.ZipFile(path) as archive:
            for member in archive.infolist():
                target = (destination / member.filename).resolve()
                if not target.is_relative_to(destination.resolve()):
                    raise ValueError('Небезопасный путь в ZIP')
                if (member.external_attr >> 16) & 0o170000 == 0o120000:
                    raise ValueError('Символическая ссылка в ZIP')
            archive.extractall(destination)
        return [item for item in destination.rglob('*') if item.is_file()]

    copied = destination / path.name
    shutil.copy2(path, copied)
    return [copied]


def _file_years(path):
    return [int(value) for value in re.findall(r'(?<!\d)(20\d{2})(?!\d)', path.name)]


def _should_read(path, years):
    if path.suffix.casefold() not in {'.csv', '.parquet', '.xlsx', '.xls'}:
        return False
    hints = _file_years(path)
    return not hints or bool(set(hints).intersection(years))


def _read_table(path, years, chunk_size):
    if path.suffix.casefold() == '.csv':
        pieces = []
        for chunk in pd.read_csv(path, sep=';', dtype=str, chunksize=chunk_size, low_memory=False):
            year_column = find_column(chunk.columns, ['год', 'year'])
            if year_column:
                parsed_year = pd.to_numeric(chunk[year_column], errors='coerce')
                chunk = chunk[parsed_year.isin(years)]
            if len(chunk):
                pieces.append(chunk)
        return pd.concat(pieces, ignore_index=True) if pieces else pd.DataFrame()

    if path.suffix.casefold() == '.parquet':
        frame = pd.read_parquet(path)
    else:
        frame = pd.read_excel(path, dtype=str)

    year_column = find_column(frame.columns, ['год', 'year'])
    if year_column:
        parsed_year = pd.to_numeric(frame[year_column], errors='coerce')
        frame = frame[parsed_year.isin(years)]
    return frame


def _harmonize(frame, indicator_name, feature_key, view, section, years, source_file, metadata_columns=()):
    if frame.empty:
        return pd.DataFrame()

    detected = {
        field: next((column for candidate in candidates for column in frame.columns
                     if text_key(column) == text_key(candidate)), None)
        for field, candidates in COLUMN_CANDIDATES.items()
    }
    if not all(detected[field] for field in ['oktmo', 'year', 'value', 'indicator_name']):
        raise ValueError(f'Не распознаны колонки ОКТМО, год, значение: {list(frame.columns)}')

    result = pd.DataFrame({
        'oktmo': frame[detected['oktmo']].map(normalize_oktmo),
        'year': pd.to_numeric(frame[detected['year']], errors='coerce'),
        'value': numeric_value(frame[detected['value']]),
    })

    if detected['indicator_name']:
        result['indicator_name'] = frame[detected['indicator_name']].map(normalize_text)
    else:
        result['indicator_name'] = indicator_name
    result['indicator_name'] = result['indicator_name'].replace('', indicator_name)

    for field in ['unit', 'municipality_name', 'municipality_type', 'region_name', 'period']:
        column = detected[field]
        result[field] = frame[column].map(normalize_text) if column else ''

    known_columns = {column for column in detected.values() if column}
    dimension_columns = [column for column in frame.columns if column not in known_columns and column not in metadata_columns]
    detail_score = np.zeros(len(frame), dtype=int)

    for column in dimension_columns:
        values = frame[column].map(text_key)
        detail_score += ((~values.isin(TOTAL_MARKERS)) & values.ne('')).astype(int).to_numpy()

    result['dimensions'] = frame[dimension_columns].fillna('').astype(str).apply(
        lambda row: repr(row.to_dict()), axis=1) if dimension_columns else ''
    result['source_comment'] = frame['comment'].fillna('') if 'comment' in frame else ''
    result['detail_score'] = detail_score
    result['feature_key'] = feature_key
    result['view'] = view
    result['requested_indicator'] = indicator_name
    result['section'] = int(section)
    result['source_file'] = Path(source_file).name

    result = result[
        result['oktmo'].notna()
        & result['year'].isin(years)
        & result['year'].mod(1).eq(0)
    ].copy()
    result['year'] = result['year'].astype(int)

    if detected['indicator_name']:
        result = result[result['indicator_name'].map(text_key).eq(text_key(indicator_name))].copy()
    # Missing and unparseable source values remain visible in the raw cache.
    result['raw_value'] = frame.loc[result.index, detected['value']].astype('string')
    result['value_missing'] = result['value'].isna()
    return result



def read_candidates(*args, **kwargs):
    raise RuntimeError('Используйте IndicatorStore.load_one для последовательной загрузки с кэшем.')


def canonicalize(raw_long, annual_periods=None):
    keys = ['oktmo', 'year', 'feature_key', 'view', 'requested_indicator']
    columns = keys + ['value', 'unit', 'municipality_name', 'municipality_type', 'region_name', 'section']
    if raw_long.empty:
        return pd.DataFrame(columns=columns), pd.DataFrame(columns=keys + ['reason'])
    data = raw_long.copy()
    data['unit'] = data['unit'].fillna('').map(normalize_text)
    # Only explicit total rows are eligible; never choose the least detailed subgroup.
    eligible = data['detail_score'].eq(0)
    annual = data['period'].map(text_key).isin({text_key(p) for p in (annual_periods or ['год', 'годовое', 'annual', 'year', 'январь-декабрь', 'значение показателя за год'])})
    rejected = data.loc[~(eligible & annual), keys].drop_duplicates()
    rejected['reason'] = 'non_total_or_nonannual'
    anomaly = data.get('source_comment', pd.Series('', index=data.index)).str.contains('аномаль', case=False, na=False)
    anomaly_rows = data.loc[anomaly, keys].drop_duplicates().assign(reason='source_anomaly')
    rejected = pd.concat([rejected, anomaly_rows], ignore_index=True)
    data = data[eligible & annual & ~anomaly]
    stats = data.groupby(keys, dropna=False).agg(
        n_values=('value', 'nunique'), n_units=('unit', 'nunique')).reset_index()
    conflicts = stats.loc[stats.n_values.gt(1) | stats.n_units.ne(1), keys].copy()
    conflicts['reason'] = 'conflicting_values_or_units'
    bad = pd.MultiIndex.from_frame(conflicts[keys])
    data = data[~pd.MultiIndex.from_frame(data[keys]).isin(bad)]
    canonical = data.groupby(keys, as_index=False, dropna=False).agg(
        value=('value', 'first'), unit=('unit', 'first'),
        municipality_name=('municipality_name', 'first'), municipality_type=('municipality_type', 'first'),
        region_name=('region_name', 'first'), section=('section', 'first'))
    return canonical.reindex(columns=columns), pd.concat([rejected, conflicts], ignore_index=True)


def load_bdmo(config, rules=FEATURE_RULES):
    raise RuntimeError('Используйте IndicatorStore.load_one из src.bdmo_workflow: явный выбор одного показателя, постоянный кэш и аудит базы.')
